"""Generate local demonstration files; existing files are never overwritten."""

import csv
from pathlib import Path

from openpyxl import Workbook


def main() -> None:
    directory = Path(__file__).parent / "generated"
    directory.mkdir(exist_ok=True)
    rows = [["Код", "Клиент", "Сумма"], ["00123", "Алматы", 1500], ["00007", "Астана", 750]]
    for filename, encoding in (("clients.csv", "utf-8-sig"), ("clients_cp1251.csv", "cp1251")):
        path = directory / filename
        if not path.exists():
            with path.open("x", encoding=encoding, newline="") as stream:
                csv.writer(stream, delimiter=";").writerows(rows)
    path = directory / "clients.xlsx"
    if not path.exists():
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Клиенты"
        for row in rows:
            sheet.append(row)
        sheet["D1"] = "Формула"
        sheet["D2"] = "=C2*2"
        workbook.create_sheet("Заказы").append(["00007", "Заказ №1"])
        try:
            with path.open("xb") as stream:
                workbook.save(stream)
        finally:
            workbook.close()
    dirty_rows = [
        ["ID", "Имя"],
        [" 001234 ", " Анна "],
        ["001234", "Анна"],
        [" ", ""],
        ["000007", ""],
    ]
    path = directory / "cleaner_demo.csv"
    if not path.exists():
        with path.open("x", encoding="utf-8-sig", newline="") as stream:
            csv.writer(stream, delimiter=";").writerows(dirty_rows)
    path = directory / "cleaner_demo.xlsx"
    if not path.exists():
        workbook = Workbook()
        workbook.active.title = "Клиенты"
        for row in dirty_rows:
            workbook.active.append(row)
        workbook.create_sheet("Справочник").append([" Не изменять ", "000042"])
        try:
            with path.open("xb") as stream:
                workbook.save(stream)
        finally:
            workbook.close()
    print(f"Примеры: {directory}")


if __name__ == "__main__":
    main()
