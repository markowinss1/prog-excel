"""Conservative XLSX support with surgical XML edits, never an openpyxl save.

Untouched ZIP members retain their uncompressed bytes. Unknown workbook features
are rejected. This is a supported subset, not a claim of arbitrary Excel fidelity.
"""

import bisect
import io
import posixpath
import re
from dataclasses import replace
from decimal import Decimal, InvalidOperation
from xml.etree import ElementTree as ET
from zipfile import ZIP_DEFLATED, ZipFile

from openpyxl import load_workbook
from openpyxl.utils.cell import coordinate_to_tuple, get_column_letter

from excel_toolkit.application.cleaning import CleaningSource
from excel_toolkit.application.models import Cell
from excel_toolkit.modules.cleaner import CleaningError, CleaningResult, Table

MAIN = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
DOC_REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
NS = {"s": MAIN}
MAX_ARCHIVE_BYTES = 128 * 1024 * 1024
MAX_EXPANDED_BYTES = 256 * 1024 * 1024
MAX_CELLS = 2_000_000

# No tables, charts, drawings, macros, external links, comments, metadata or extensions.
PART = re.compile(
    r"(?:\[Content_Types\]\.xml|_rels/\.rels|docProps/(?:app|core)\.xml|"
    r"xl/(?:workbook\.xml|_rels/workbook\.xml\.rels|styles\.xml|sharedStrings\.xml|"
    r"theme/theme[0-9]+\.xml|worksheets/sheet[0-9]+\.xml))\Z"
)
SHEET_TAGS = set(
    "worksheet sheetPr outlinePr pageSetUpPr dimension sheetViews sheetView pane selection "
    "sheetFormatPr cols col sheetData row c v f is t pageMargins".split()
)
BOOK_TAGS = set(
    "workbook fileVersion workbookPr bookViews workbookView sheets sheet "
    "definedNames calcPr workbookProtection".split()
)
FEATURE_BAG_ERROR = (
    "Неподдерживаемая часть XLSX: featurePropertyBag. Она может содержать свойства "
    "элементов управления и ссылки из стилей. Проверка этих связей пока не реализована; "
    "безопасность очистки не гарантируется. Исходный файл не изменён."
)
FEATURE_BAG_NS = "http://schemas.microsoft.com/office/spreadsheetml/2022/featurepropertybag"
FEATURE_BAG_CONTENT_TYPE = "application/vnd.ms-excel.featurepropertybag+xml"
FEATURE_BAG_STYLE_URIS = {
    "{C7286773-470A-42A8-94C5-96B5CB345126}",  # xfComplement
    "{0417FA29-78FA-4A13-93AC-8FF0FAFDF519}",  # DXFComplement
}


def parse_xml(data: bytes) -> ET.Element:
    # Reject DTD/entity inputs before either XML parser sees them.
    text = data.decode("utf-8-sig")
    if "\x00" in text:
        raise CleaningError("Поддерживается только UTF-8 XML внутри XLSX.")
    if "<!DOCTYPE" in text.upper() or "<!ENTITY" in text.upper():
        raise CleaningError("XLSX с DTD или XML entities не поддерживается.")
    return ET.fromstring(data)


def _check_tags(root: ET.Element, allowed: set[str]) -> None:
    qualified = {f"{{{MAIN}}}{name}" for name in allowed}
    for node in root.iter():
        if node.tag not in qualified:
            raise CleaningError(f"Неподдерживаемая структура XLSX: {node.tag.rsplit('}', 1)[-1]}.")
        if node.tag == f"{{{MAIN}}}t" and node.text and re.search(r"_x[0-9a-fA-F]{4}_", node.text):
            raise CleaningError("OOXML escape-последовательности в тексте пока не поддерживаются.")
        if any(
            key.startswith("{") and key != "{http://www.w3.org/XML/1998/namespace}space"
            for key in node.attrib
        ):
            # Workbook sheet relationship identifiers are the sole needed exception.
            if node.tag != f"{{{MAIN}}}sheet" or set(node.attrib) - {
                "name",
                "sheetId",
                "state",
                f"{{{DOC_REL}}}id",
            }:
                raise CleaningError("Расширенные атрибуты XLSX не поддерживаются.")


def _check_style_extensions(styles: ET.Element) -> None:
    """Do not silently accept implicit bag references when its part/rels are absent.

    Styles are copied verbatim, but extensions can change what an empty cell means
    (e.g. a checkbox with a default state). Copying bytes alone cannot make row
    deletion safe. Ordinary SpreadsheetML formatting remains supported.
    """
    for node in styles.iter():
        if node.tag.startswith(f"{{{FEATURE_BAG_NS}}}") or node.get("uri", "").upper() in (
            FEATURE_BAG_STYLE_URIS
        ):
            raise CleaningError(FEATURE_BAG_ERROR)
    for node in styles.iter():
        if (
            not node.tag.startswith(f"{{{MAIN}}}")
            or node.tag in {f"{{{MAIN}}}ext", f"{{{MAIN}}}extLst"}
            or any(key.startswith("{") for key in node.attrib)
        ):
            raise CleaningError("Неподдерживаемые расширения стилей XLSX.")


def _view_error(detail: str) -> CleaningError:
    return CleaningError(f"Неподдерживаемая структура XLSX: pane/sheetViews — {detail}.")


def _view_cell(reference: str) -> tuple[int, int]:
    if not re.fullmatch(r"\$?[A-Z]{1,3}\$?[1-9][0-9]{0,6}", reference):
        raise _view_error(f"некорректная координата {reference!r}")
    row, column = coordinate_to_tuple(reference.replace("$", ""))
    if row > 1048576 or column > 16384:
        raise _view_error("координата вне границ листа")
    return row, column


def _split_count(pane: ET.Element, name: str, limit: int) -> int:
    try:
        value = Decimal(pane.get(name, "0"))
    except InvalidOperation as exc:
        raise _view_error(f"неверное значение {name}") from exc
    if not value.is_finite() or not 0 <= value < limit or value != int(value):
        raise _view_error(f"{name} должен задавать целое число закреплённых строк/столбцов")
    return int(value)


def _check_sheet_views(root: ET.Element) -> ET.Element | None:
    views = root.findall("s:sheetViews", NS)
    # The allowlist alone does not enforce nesting or cardinality.
    for tag, path in (
        ("sheetViews", "s:sheetViews"),
        ("sheetView", "s:sheetViews/s:sheetView"),
        ("pane", "s:sheetViews/s:sheetView/s:pane"),
        ("selection", "s:sheetViews/s:sheetView/s:selection"),
    ):
        if len(list(root.iter(f"{{{MAIN}}}{tag}"))) != len(root.findall(path, NS)):
            raise _view_error(f"неверное расположение {tag}")
    if len(views) > 1:
        raise _view_error("несколько sheetViews")
    if not views:
        return None
    views = views[0]
    if views.attrib or not len(views):
        raise _view_error("неверный sheetViews")
    for view in views:
        if view.tag != f"{{{MAIN}}}sheetView":
            raise _view_error("ожидается sheetView")
        panes = view.findall("s:pane", NS)
        selections = view.findall("s:selection", NS)
        if len(panes) > 1 or len(selections) > 4 or len(view) != len(panes) + len(selections):
            raise _view_error("неверный состав sheetView")
        if "topLeftCell" in view.attrib:
            _view_cell(view.attrib["topLeftCell"])
        available = {"topLeft"}
        if panes:
            pane = panes[0]
            if (
                view[0] is not pane
                or len(pane)
                or set(pane.attrib) - {"xSplit", "ySplit", "topLeftCell", "activePane", "state"}
            ):
                raise _view_error("неизвестные свойства или расположение pane")
            if pane.get("state") != "frozen":
                raise _view_error("поддерживается только state=frozen; split/frozenSplit запрещены")
            if (
                view.get("rightToLeft", "0") not in {"0", "false"}
                or view.get("view", "normal") != "normal"
            ):
                raise _view_error("закрепление поддерживается в обычном виде слева направо")
            x = _split_count(pane, "xSplit", 16384)
            y = _split_count(pane, "ySplit", 1048576)
            row, column = _view_cell(pane.get("topLeftCell", ""))
            if not (x or y) or row <= y or column <= x:
                raise _view_error("topLeftCell не согласован с границей закрепления")
            if x:
                available.add("topRight")
            if y:
                available.add("bottomLeft")
            if x and y:
                available.add("bottomRight")
            if pane.get("activePane", "topLeft") not in available:
                raise _view_error("activePane не соответствует закреплению")
        for selection in selections:
            if len(selection) or set(selection.attrib) - {
                "pane",
                "activeCell",
                "activeCellId",
                "sqref",
            }:
                raise _view_error("неизвестные свойства selection")
            if selection.get("pane", "topLeft") not in available:
                raise _view_error("selection ссылается на отсутствующую область")
            if "activeCell" in selection.attrib:
                _view_cell(selection.attrib["activeCell"])
            ranges = selection.get("sqref", "A1").split()
            if not ranges:
                raise _view_error("пустой sqref")
            for reference in ranges:
                ends = reference.split(":")
                if len(ends) > 2:
                    raise _view_error("неверный диапазон selection")
                start, end = _view_cell(ends[0]), _view_cell(ends[-1])
                if start[0] > end[0] or start[1] > end[1]:
                    raise _view_error("обратный диапазон selection")
            if not 0 <= int(selection.get("activeCellId", "0")) < len(ranges):
                raise _view_error("activeCellId вне списка диапазонов")
    return views


def _remap_sheet_views(views: ET.Element | None, removed: list[int]) -> None:
    """Keep freeze geometry; shift surviving view references with their source rows.

    A deleted anchor or any deletion inside a selection is deliberately rejected:
    choosing a replacement active cell/range would change the user's view semantics.
    """
    if views is None or not removed:
        return

    def shift(reference: str) -> str:
        row, _ = _view_cell(reference)
        offset = bisect.bisect_left(removed, row)
        if offset < len(removed) and removed[offset] == row:
            raise CleaningError(
                "Удаление строк заблокировано: удаляется координата прокрутки/выделения "
                "sheetViews. Переместите выделение/прокрутку в Excel или отключите удаление строк."
            )
        return re.sub(r"[0-9]+$", str(row - offset), reference)

    for view in views:
        pane = view.find("s:pane", NS)
        if pane is not None and removed[0] <= _split_count(pane, "ySplit", 1048576):
            raise CleaningError(
                "Удаление строк заблокировано: строка входит в закреплённую область pane. "
                "Отключите удаление строк или измените закрепление в Excel."
            )
        for node in view.iter():
            for attribute in ("topLeftCell", "activeCell"):
                if attribute in node.attrib:
                    node.set(attribute, shift(node.attrib[attribute]))
            if "sqref" in node.attrib:
                references = []
                for reference in node.attrib["sqref"].split():
                    ends = reference.split(":")
                    first, last = _view_cell(ends[0])[0], _view_cell(ends[-1])[0]
                    if bisect.bisect_right(removed, last) != bisect.bisect_left(removed, first):
                        raise CleaningError(
                            "Удаление строк заблокировано: затронут диапазон выделения "
                            "selection/sqref. Измените выделение в Excel "
                            "или отключите удаление строк."
                        )
                    references.append(":".join(shift(end) for end in ends))
                node.set("sqref", " ".join(references))


def inspect_archive(snapshot: bytes) -> tuple[dict[str, str], bool, dict[str, bytes]]:
    with ZipFile(io.BytesIO(snapshot)) as archive:
        infos = archive.infolist()
        names = [item.filename for item in infos]
        if len(set(names)) != len(names) or len(names) > 1000:
            raise CleaningError("Некорректный или слишком большой ZIP-каталог XLSX.")
        if sum(item.file_size for item in infos) > MAX_EXPANDED_BYTES:
            raise CleaningError("Распакованная книга превышает лимит 256 MiB.")
        for name in names:
            if name.startswith("xl/featurePropertyBag/"):
                raise CleaningError(FEATURE_BAG_ERROR)
            if not PART.fullmatch(name):
                raise CleaningError(f"Неподдерживаемая часть XLSX: {name}.")
        roots = {name: parse_xml(archive.read(name)) for name in names}
        for declaration in roots["[Content_Types].xml"]:
            if declaration.get("ContentType") == FEATURE_BAG_CONTENT_TYPE:
                raise CleaningError(FEATURE_BAG_ERROR)
        for name, root in roots.items():
            if name.endswith(".rels"):
                for relation in root:
                    if relation.get("TargetMode") == "External":
                        raise CleaningError("Внешние ссылки XLSX не поддерживаются.")
                    kind = relation.get("Type", "").rsplit("/", 1)[-1]
                    if kind == "FeaturePropertyBag":
                        raise CleaningError(FEATURE_BAG_ERROR)
                    if kind not in {
                        "officeDocument",
                        "core-properties",
                        "extended-properties",
                        "worksheet",
                        "styles",
                        "theme",
                        "sharedStrings",
                    }:
                        raise CleaningError(f"Неподдерживаемая связь XLSX: {kind}.")
        styles = roots.get("xl/styles.xml")
        if styles is not None:
            _check_style_extensions(styles)
        book = roots["xl/workbook.xml"]
        _check_tags(book, BOOK_TAGS)  # Nonempty definedNames introduces forbidden definedName.
        protection = book.find("s:workbookProtection", NS)
        if protection is not None and protection.attrib:
            raise CleaningError("Неподдерживаемая структура XLSX: защита книги.")
        strings = roots.get("xl/sharedStrings.xml")
        if strings is not None:
            _check_tags(strings, {"sst", "si", "t"})  # Rich text must not be flattened.
        relationships = {}
        for relation in roots["xl/_rels/workbook.xml.rels"]:
            target = relation.attrib["Target"]
            part = posixpath.normpath(posixpath.join("xl", target))
            if target.startswith("/"):
                part = target.lstrip("/")
            relationships[relation.attrib["Id"]] = part
        sheets = {}
        for node in book.findall("s:sheets/s:sheet", NS):
            part = relationships[node.attrib[f"{{{DOC_REL}}}id"]]
            if not re.fullmatch(r"xl/worksheets/sheet[0-9]+\.xml", part):
                raise CleaningError("Неподдерживаемый тип листа XLSX.")
            if node.attrib["name"] in sheets or part in sheets.values():
                raise CleaningError("Повторяющиеся листы XLSX не поддерживаются.")
            sheets[node.attrib["name"]] = part
        formulas = False
        sheet_views = {}
        for part in sheets.values():
            root = roots[part]
            _check_tags(root, SHEET_TAGS)
            views = _check_sheet_views(root)
            if views is not None:
                sheet_views[part] = ET.tostring(views, encoding="utf-8")
            if root.tag != f"{{{MAIN}}}worksheet" or len(root.findall("s:sheetData", NS)) != 1:
                raise CleaningError("Некорректная структура листа XLSX.")
            if any(int(node.get("outlineLevel", "0")) for node in root.iter()):
                raise CleaningError("Неподдерживаемая структура XLSX: группировка строк/столбцов.")
            previous_row = 0
            for row in root.findall("s:sheetData/s:row", NS):
                index = int(row.attrib["r"])
                if index <= previous_row or index > 1048576:
                    raise CleaningError("Некорректный порядок строк XLSX.")
                previous_row = index
                previous_column = 0
                for cell in row.findall("s:c", NS):
                    r, c = coordinate_to_tuple(cell.attrib["r"])
                    if r != index or c <= previous_column or c > 16384:
                        raise CleaningError("Некорректные координаты ячеек XLSX.")
                    previous_column = c
                    if set(cell.attrib) - {"r", "s", "t"}:
                        raise CleaningError("Метаданные ячеек XLSX не поддерживаются.")
                    formula = cell.find("s:f", NS)
                    if formula is not None:
                        if formula.attrib:
                            raise CleaningError(
                                "Общие, массивные и табличные формулы не поддерживаются."
                            )
                        formulas = True
        return sheets, formulas, sheet_views


def read_table(snapshot: bytes, selected: str, part: str) -> tuple[Table, int]:
    """Return the logical rectangle and count of physical rows beyond it.

    A stored v/is/f payload anchors the rectangle, even if its value is blank.
    Bare rows/cells, styles and dimension metadata alone do not extend it.
    Rows outside this rectangle are preserved, not reported as removed.
    """
    with ZipFile(io.BytesIO(snapshot)) as archive:
        root = parse_xml(archive.read(part))
    last_row = last_column = 0
    for cell in root.findall("s:sheetData/s:row/s:c", NS):
        if (
            cell.find("s:v", NS) is not None
            or cell.find("s:is", NS) is not None
            or (cell.find("s:f", NS) is not None)
        ):
            r, c = coordinate_to_tuple(cell.attrib["r"])
            last_row, last_column = max(last_row, r), max(last_column, c)
    if last_row * last_column > MAX_CELLS:
        raise CleaningError("Область данных XLSX превышает лимит 2 000 000 ячеек.")
    trailing_rows = sum(
        int(row.attrib["r"]) > last_row for row in root.findall("s:sheetData/s:row", NS)
    )
    if not last_row:
        return Table(()), trailing_rows
    with io.BytesIO(snapshot) as stream:
        book = load_workbook(stream, read_only=True, data_only=False, keep_links=False)
        try:
            sheet = book[selected]
            sheet.reset_dimensions()
            table = Table(
                tuple(
                    tuple(Cell(c.value, c.data_type, c.number_format) for c in row)
                    for row in sheet.iter_rows(max_row=last_row, max_col=last_column)
                )
            )
        finally:
            book.close()
    # Excel serials 59 and 60 can decode to the same date; time decoding also rounds.
    # Never allow lossy reader normalization to turn distinct stored values into duplicates.
    decoded_numbers = {}
    for cell in root.findall("s:sheetData/s:row/s:c", NS):
        value = cell.find("s:v", NS)
        if cell.get("t", "n") != "n" or value is None or value.text is None:
            continue
        if cell.find("s:f", NS) is not None:
            continue
        row, column = coordinate_to_tuple(cell.attrib["r"])
        decoded = table.rows[row - 1][column - 1].value
        raw = Decimal(value.text)
        key = type(decoded), decoded
        if not raw.is_finite() or (key in decoded_numbers and decoded_numbers[key] != raw):
            raise CleaningError(
                "Разные числовые значения/даты XLSX неразличимы после чтения; "
                "безопасное сравнение не гарантируется."
            )
        decoded_numbers[key] = raw
    return table, trailing_rows


def validate_result(source: CleaningSource, result: CleaningResult) -> None:
    rows = result.source_rows
    if (
        len(rows) != len(result.table.rows)
        or any(type(r) is not int or r < 1 or r > len(source.table.rows) for r in rows)
        or any(a >= b for a, b in zip(rows, rows[1:], strict=False))
    ):
        raise CleaningError("Результат не соответствует исходным строкам.")
    if source.contains_formulas and len(rows) != len(source.table.rows):
        raise CleaningError(
            "Удаление строк заблокировано: в книге есть формулы, включая другие листы."
        )
    if source.xlsx_sheet_views is not None and len(rows) != len(source.table.rows):
        removed = sorted(set(range(1, len(source.table.rows) + 1)) - set(rows))
        _remap_sheet_views(parse_xml(source.xlsx_sheet_views), removed)
    for index, row in zip(rows, result.table.rows, strict=True):
        original = source.table.rows[index - 1]
        if len(original) != len(row):
            raise CleaningError("Структура результата не соответствует оригиналу.")
        for before, after in zip(original, row, strict=True):
            allowed = before
            if isinstance(before.value, str) and before.data_type != "f":
                allowed = replace(before, value=before.value.strip())
            if after not in (before, allowed) or type(after.value) is not type(before.value):
                raise CleaningError("Экспорт допускает только изменения Cleaner.")


def write_xlsx(source: CleaningSource, result: CleaningResult, stream) -> None:
    if source.xlsx_snapshot is None or source.sheet_part is None:
        raise CleaningError("Для XLSX требуется исходный снимок книги.")
    with ZipFile(io.BytesIO(source.xlsx_snapshot)) as archive:
        removed = sorted(set(range(1, len(source.table.rows) + 1)) - set(result.source_rows))
        changed = {}
        for source_index, row in zip(result.source_rows, result.table.rows, strict=True):
            for col, (old, new) in enumerate(
                zip(source.table.rows[source_index - 1], row, strict=True), start=1
            ):
                if old != new:
                    changed[source_index, col] = new.value
        modified_xml = None
        if removed or changed:
            root = parse_xml(archive.read(source.sheet_part))
            _remap_sheet_views(root.find("s:sheetViews", NS), removed)
            data = root.find("s:sheetData", NS)
            removed_set = set(removed)
            retained = []
            for row in data:
                old_row = int(row.attrib["r"])
                if old_row in removed_set:
                    continue
                retained.append(row)
                new_row = old_row - bisect.bisect_left(removed, old_row)
                row.set("r", str(new_row))
                for cell in row.findall("s:c", NS):
                    _, col = coordinate_to_tuple(cell.attrib["r"])
                    cell.set("r", f"{get_column_letter(col)}{new_row}")
                    if (old_row, col) in changed:
                        # Inline strings prevent =... text from becoming a formula.
                        for child in list(cell):
                            cell.remove(child)
                        cell.set("t", "inlineStr")
                        inline = ET.SubElement(cell, f"{{{MAIN}}}is")
                        text = ET.SubElement(inline, f"{{{MAIN}}}t")
                        text.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
                        text.text = changed[old_row, col]
            data[:] = retained
            dimension = root.find("s:dimension", NS)
            if dimension is not None:
                # Let readers compute dimensions; stale metadata must not truncate data.
                root.remove(dimension)
            # ElementTree emits literal CR in text; XML readers would normalize it to LF.
            modified_xml = ET.tostring(root, encoding="utf-8", xml_declaration=True).replace(
                b"\r", b"&#13;"
            )
        with ZipFile(stream, "w", compression=ZIP_DEFLATED) as output:
            output.comment = archive.comment
            for info in archive.infolist():
                payload = (
                    modified_xml
                    if modified_xml is not None and info.filename == source.sheet_part
                    else archive.read(info.filename)
                )
                output.writestr(info, payload)
