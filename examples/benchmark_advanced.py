"""Measure all/selected-column cleaning and export on 10,000 synthetic records."""

import argparse
import csv
import hashlib
import json
import platform
import statistics
import tempfile
from pathlib import Path
from time import perf_counter

import openpyxl
from openpyxl import Workbook, load_workbook

from excel_toolkit.application.cleaning import CleaningService
from excel_toolkit.infrastructure.files.cleaning import FileCleaningStorage
from excel_toolkit.modules.cleaner import CleaningPlan


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--repeats", type=int, default=3)
    args = parser.parse_args()
    if args.repeats < 1:
        parser.error("--repeats must be positive")
    rows = [["ID", "Name", "City"]]
    rows.extend(
        [f"{i // 2:06d}", f"Client {i // 2}", "Astana" if i % 2 == 0 else "Almaty"]
        for i in range(9900)
    )
    rows.extend([["", "Review", "Astana"] for _ in range(100)])
    results = []
    service = CleaningService(FileCleaningStorage())
    with tempfile.TemporaryDirectory(prefix="cleaner-advanced-benchmark-") as directory:
        for suffix in (".csv", ".xlsx"):
            path = Path(directory) / f"source{suffix}"
            if suffix == ".csv":
                with path.open("x", encoding="utf-8-sig", newline="") as stream:
                    csv.writer(stream).writerows(rows)
            else:
                book = Workbook()
                for row in rows:
                    book.active.append(row)
                book.active.freeze_panes = "A2"
                book.create_sheet("Other").append([" unchanged ", "00001"])
                book.save(path)
                book.close()
            before = hashlib.sha256(path.read_bytes()).hexdigest(), path.stat().st_mtime_ns
            for mode, columns in (("all", None), ("selected_id", (0,))):
                samples = []
                for repetition in range(args.repeats + 1):
                    output = Path(directory) / f"{mode}-{repetition}{suffix}"
                    start = perf_counter()
                    job = service.analyze(
                        path, CleaningPlan(True, True, True, True, duplicate_columns=columns)
                    )
                    analyzed = perf_counter()
                    service.export(job, output)
                    exported = perf_counter()
                    if repetition:
                        samples.append({"analyze": analyzed - start, "export": exported - analyzed})
                    report = job.result.report
                    expected = (9902, 99, 0, 0) if columns is None else (10001, 0, 4950, 100)
                    assert (
                        report.final_rows,
                        report.removed_duplicates,
                        report.conflicting_duplicate_groups,
                        report.missing_key_rows,
                    ) == expected
                # Read-back and assertions do not contribute to timing.
                restored = service.analyze(output, CleaningPlan())
                assert restored.result.table == job.result.table
                if suffix == ".xlsx":
                    book = load_workbook(output)
                    try:
                        assert book.active.freeze_panes == "A2"
                        assert book["Other"]["A1"].value == " unchanged "
                    finally:
                        book.close()
                assert before == (
                    hashlib.sha256(path.read_bytes()).hexdigest(),
                    path.stat().st_mtime_ns,
                )
                results.append(
                    {
                        "format": suffix,
                        "mode": mode,
                        "report": {
                            name: getattr(report, name)
                            for name in (
                                "original_rows",
                                "final_rows",
                                "removed_duplicates",
                                "conflicting_duplicate_groups",
                                "missing_key_rows",
                            )
                        },
                        "samples_seconds": samples,
                        "median_seconds": {
                            key: statistics.median(sample[key] for sample in samples)
                            for key in ("analyze", "export")
                        },
                    }
                )
    print(
        json.dumps(
            {
                "python": platform.python_version(),
                "openpyxl": openpyxl.__version__,
                "platform": platform.platform(),
                "records": 10_000,
                "warmups": 1,
                "repeats": args.repeats,
                "results": results,
            },
            indent=2,
        )
    )


if __name__ == "__main__":
    main()
