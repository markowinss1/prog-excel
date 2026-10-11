from datetime import datetime
from pathlib import Path

import pytest
from openpyxl import Workbook

from excel_toolkit.application.preview import PreviewService
from excel_toolkit.infrastructure.files.reader import FilePreviewReader


@pytest.fixture
def service() -> PreviewService:
    return PreviewService(FilePreviewReader())


@pytest.fixture
def xlsx_file(tmp_path: Path) -> Path:
    path = tmp_path / "Клиенты.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Клиенты"
    sheet.append(["Код", "Имя", "Количество", "Формула", "Дата", "Флаг"])
    sheet.append(["00123", "Алматы", 12, "=C2*2", datetime(2026, 1, 2), True])
    sheet["C2"].number_format = "00000"
    second = workbook.create_sheet("Заказы")
    second.append(["00007", "Заказ"])
    workbook.create_sheet("Пустой")
    workbook.save(path)
    workbook.close()
    return path
