import hashlib
import posixpath
from dataclasses import replace
from datetime import datetime
from xml.etree import ElementTree as ET
from zipfile import ZipFile

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.worksheet.views import Selection, SheetView

from excel_toolkit.application.cleaning import CleaningService
from excel_toolkit.infrastructure.files.cleaning import FileCleaningStorage
from excel_toolkit.infrastructure.files.cleaning_xlsx import NS, inspect_archive
from excel_toolkit.modules.cleaner import CleaningError, CleaningPlan, clean


@pytest.fixture
def cleaner():
    return CleaningService(FileCleaningStorage())


def fingerprint(path):
    return hashlib.sha256(path.read_bytes()).digest(), path.stat().st_mtime_ns


def make_book(path, freeze="A2", configure=None):
    book = Workbook()
    sheet = book.active
    sheet.title = "Data"
    sheet.append([" ID ", "Value", "Date"])
    sheet.append([" 00001 ", 12, datetime(2026, 1, 2)])
    sheet.append([" 00002 ", 24, "2026-01-02"])
    sheet.append([" 00003 ", 36, None])
    sheet.append([None, " "])
    sheet.append(["00003", 36, None])
    sheet.append([" 00004 ", 48, False])
    sheet["A7"].font = Font(bold=True)
    sheet["B7"].number_format = "00000"
    sheet.row_dimensions[7].height = 32
    sheet.freeze_panes = freeze
    other = book.create_sheet("Other")
    other.append([" untouched ", "00009"])
    other.freeze_panes = "C4"
    if configure:
        configure(book)
    book.save(path)
    book.close()


def rewrite(path, transform):
    with ZipFile(path) as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}
    transform(parts)
    with ZipFile(path, "w") as archive:
        for name, payload in parts.items():
            archive.writestr(name, payload)


def mutate_sheet(path, transform):
    def change(parts):
        name = "xl/worksheets/sheet1.xml"
        root = ET.fromstring(parts[name])
        transform(root)
        parts[name] = ET.tostring(root)

    rewrite(path, change)


def check_package(path):
    """CRC, XML well-formedness, OPC targets and the application's supported profile."""
    with ZipFile(path) as archive:
        assert archive.testzip() is None
        for name in archive.namelist():
            root = ET.fromstring(archive.read(name))
            if name.endswith(".rels"):
                base = posixpath.dirname(posixpath.dirname(name))
                for relation in root:
                    target = relation.attrib["Target"]
                    target = (
                        target.lstrip("/")
                        if target.startswith("/")
                        else posixpath.normpath(posixpath.join(base, target))
                    )
                    assert target in archive.namelist()
            if name == "[Content_Types].xml":
                for node in root:
                    if "PartName" in node.attrib:
                        assert node.attrib["PartName"].lstrip("/") in archive.namelist()
    inspect_archive(path.read_bytes())


@pytest.mark.parametrize("freeze", [None, "A2", "B1", "A4", "C4"])
@pytest.mark.parametrize("plan", [CleaningPlan(), CleaningPlan(True, True, True, True)])
def test_freeze_round_trip_and_data_integrity(cleaner, tmp_path, freeze, plan):
    path = tmp_path / "source.xlsx"
    make_book(path, freeze)
    before = fingerprint(path)
    job = cleaner.analyze(path, plan)
    output = tmp_path / "out.xlsx"
    cleaner.export(job, output)
    assert fingerprint(path) == before
    check_package(output)
    with ZipFile(path) as original, ZipFile(output) as result:
        assert original.namelist() == result.namelist()
        for name in original.namelist():
            if name != job.source.sheet_part or not plan.trim_whitespace:
                assert original.read(name) == result.read(name)
        a = ET.fromstring(original.read(job.source.sheet_part)).find("s:sheetViews", NS)
        b = ET.fromstring(result.read(job.source.sheet_part)).find("s:sheetViews", NS)
        assert ET.tostring(a) == ET.tostring(b)
    book = load_workbook(output)
    try:
        sheet = book["Data"]
        assert sheet.freeze_panes == freeze
        assert book["Other"].freeze_panes == "C4"
        assert book["Other"]["A1"].value == " untouched "
        assert sheet["A2"].value == ("00001" if plan.trim_whitespace else " 00001 ")
        assert sheet["A2"].data_type == "s"
        assert sheet["B2"].value == 12 and sheet["B2"].data_type == "n"
        assert sheet["C2"].value == datetime(2026, 1, 2)
        assert sheet["C3"].value == "2026-01-02"
        last = 5 if plan.remove_duplicates else 7
        assert sheet.max_row == last
        assert sheet.cell(last, 1).font.bold
        assert sheet.cell(last, 2).number_format == "00000"
        assert sheet.row_dimensions[last].height == 32
        if plan.remove_duplicates:
            assert [sheet.cell(r, 1).value for r in range(1, last + 1)] == [
                "ID",
                "00001",
                "00002",
                "00003",
                "00004",
            ]
            assert job.result.report.removed_empty_rows == 1
            assert job.result.report.removed_duplicates == 1
    finally:
        book.close()
    assert cleaner.analyze(output, CleaningPlan()).result.table == job.result.table


def test_scrolled_pane_and_multirange_selection_follow_surviving_rows(cleaner, tmp_path):
    path = tmp_path / "scrolled.xlsx"

    def configure(book):
        view = book.active.sheet_view
        view.topLeftCell = "A7"
        view.pane.topLeftCell = "C7"
        view.selection[-1] = Selection(
            pane="bottomRight", activeCell="D9", activeCellId=1, sqref="C7:C8 D9:E10"
        )
        # Additional workbook windows must be preserved and remapped too.
        book.views.append(book.views[0])
        book.active.views.sheetView.append(
            SheetView(
                workbookViewId=1,
                topLeftCell="A8",
                selection=[Selection(activeCell="A8", sqref="A8")],
            )
        )

    make_book(path, "C4", configure)
    job = cleaner.analyze(path, CleaningPlan(True, True, True, True))
    output = tmp_path / "shifted.xlsx"
    cleaner.export(job, output)
    check_package(output)
    book = load_workbook(output)
    try:
        view = book.active.sheet_view
        assert view.topLeftCell == "A5"
        assert (view.pane.xSplit, view.pane.ySplit, view.pane.topLeftCell) == (2, 3, "C5")
        assert view.selection[-1].activeCell == "D7"
        assert view.selection[-1].sqref == "C5:C6 D7:E8"
        assert view.selection[-1].activeCellId == 1
        assert book.active.views.sheetView[1].topLeftCell == "A6"
    finally:
        book.close()


@pytest.mark.parametrize("operation", ["empty", "duplicate", "no_header"])
def test_deletion_inside_frozen_rows_is_rejected_at_analysis_and_export(
    cleaner, tmp_path, operation
):
    path = tmp_path / "boundary.xlsx"

    def configure(book):
        sheet = book.active
        if operation == "duplicate":
            for column in range(1, 4):
                sheet.cell(3, column).value = sheet.cell(2, column).value
        else:
            for cell in sheet[1 if operation == "no_header" else 3]:
                cell.value = None

    make_book(path, "C4", configure)
    before = fingerprint(path)
    plan = CleaningPlan(True, True, True, True, has_header=operation != "no_header")
    with pytest.raises(CleaningError, match="закреплённую область"):
        cleaner.analyze(path, plan)
    # Export revalidates even if a caller supplies an independently computed result.
    job = cleaner.analyze(path, CleaningPlan(trim_whitespace=True))
    job = replace(job, result=clean(job.source.table, plan))
    with pytest.raises(CleaningError, match="закреплённую область"):
        cleaner.export(job, tmp_path / "out.xlsx")
    assert not (tmp_path / "out.xlsx").exists()
    assert not list(tmp_path.glob(".cleaner-*"))
    assert fingerprint(path) == before


@pytest.mark.parametrize("kind", ["viewport", "active", "range", "view_top_left"])
def test_deleted_view_anchors_and_intersecting_selections_are_rejected(cleaner, tmp_path, kind):
    path = tmp_path / "anchors.xlsx"

    def configure(book):
        view = book.active.sheet_view
        if kind == "viewport":
            view.pane.topLeftCell = "A5"
        elif kind == "view_top_left":
            view.topLeftCell = "A5"
        elif kind == "active":
            view.selection[0].activeCell = "A5"
            view.selection[0].sqref = "A5"
        else:
            view.selection[0].activeCell = "A4"
            view.selection[0].sqref = "A4:A7"

    make_book(path, configure=configure)
    with pytest.raises(CleaningError, match="Удаление строк заблокировано"):
        cleaner.analyze(path, CleaningPlan(remove_empty_rows=True))
    job = cleaner.analyze(path, CleaningPlan(trim_whitespace=True))
    cleaner.export(job, tmp_path / "trim.xlsx")
    check_package(tmp_path / "trim.xlsx")


@pytest.mark.parametrize(
    "attribute,value",
    [
        ("state", "split"),
        ("state", "frozenSplit"),
        ("state", "unknown"),
        ("ySplit", "0.5"),
        ("ySplit", "-1"),
        ("ySplit", "NaN"),
        ("ySplit", "INF"),
        ("ySplit", "bad"),
        ("ySplit", "1048576"),
        ("xSplit", "16384"),
        ("topLeftCell", "A1"),
        ("topLeftCell", "A1048577"),
        ("topLeftCell", "XFE2"),
        ("topLeftCell", "A0"),
        ("activePane", "topRight"),
        ("unknown", "1"),
    ],
)
def test_unsupported_pane_has_safe_refusal(cleaner, tmp_path, attribute, value):
    path = tmp_path / "unsupported.xlsx"
    make_book(path)
    mutate_sheet(
        path, lambda root: root.find("s:sheetViews/s:sheetView/s:pane", NS).set(attribute, value)
    )
    before = fingerprint(path)
    with pytest.raises(CleaningError, match="pane/sheetViews"):
        cleaner.analyze(path, CleaningPlan())
    assert fingerprint(path) == before


@pytest.mark.parametrize("kind", ["duplicate", "misplaced", "right_to_left", "selection", "sqref"])
def test_invalid_view_structure_is_rejected(cleaner, tmp_path, kind):
    path = tmp_path / "structure.xlsx"
    make_book(path)

    def change(root):
        view = root.find("s:sheetViews/s:sheetView", NS)
        pane = view.find("s:pane", NS)
        if kind == "duplicate":
            view.append(pane)
        elif kind == "misplaced":
            view.remove(pane)
            root.append(pane)
        elif kind == "right_to_left":
            view.set("rightToLeft", "1")
        elif kind == "selection":
            view.find("s:selection", NS).set("pane", "bottomRight")
        else:
            view.find("s:selection", NS).set("sqref", "A7:A2")

    mutate_sheet(path, change)
    with pytest.raises(CleaningError, match="pane/sheetViews"):
        cleaner.analyze(path, CleaningPlan())


def test_formulas_with_freeze_preserve_text_and_cached_results(cleaner, tmp_path):
    path = tmp_path / "formula.xlsx"

    def configure(book):
        book.active["D2"] = "=B2*2"

    make_book(path, configure=configure)

    def cached(root):
        root.find('s:sheetData/s:row/s:c[@r="D2"]/s:v', NS).text = "24"

    mutate_sheet(path, cached)
    with pytest.raises(CleaningError, match="формулы"):
        cleaner.analyze(path, CleaningPlan(True, True, True, True))
    job = cleaner.analyze(path, CleaningPlan(trim_whitespace=True))
    output = tmp_path / "out.xlsx"
    cleaner.export(job, output)
    assert job.result.report.warnings
    for data_only, value in ((False, "=B2*2"), (True, 24)):
        book = load_workbook(output, data_only=data_only)
        try:
            assert book.active["D2"].value == value
            assert book.active.freeze_panes == "A2"
        finally:
            book.close()


@pytest.mark.parametrize("kind", ["checkbox", "empty", "unknown", "external", "dangling"])
def test_feature_property_bags_remain_explicitly_unsupported(cleaner, feature_bag_xlsx, kind):
    path = feature_bag_xlsx

    def change(parts):
        name = "xl/featurePropertyBag/featurePropertyBag.xml"
        if kind == "empty":
            root = ET.fromstring(parts[name])
            root.clear()
            parts[name] = ET.tostring(root)
        elif kind == "unknown":
            parts[name] = parts[name].replace(b'"Checkbox"', b'"UnknownFeature"')
        elif kind == "dangling":
            parts[name] = parts[name].replace(b">0</bagId>", b">999</bagId>")
        elif kind == "external":
            parts["xl/featurePropertyBag/_rels/featurePropertyBag.xml.rels"] = (
                b'<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/'
                b'relationships"><Relationship Id="external" Type="urn:unknown" '
                b'Target="https://example.invalid/metadata" TargetMode="External"/></Relationships>'
            )

    rewrite(path, change)
    before = fingerprint(path)
    with pytest.raises(CleaningError, match="featurePropertyBag.*Проверка этих связей"):
        cleaner.analyze(path, CleaningPlan(trim_whitespace=True))
    assert fingerprint(path) == before


def test_feature_bag_relationship_without_part_is_rejected(cleaner, feature_bag_xlsx):
    rewrite(
        feature_bag_xlsx, lambda parts: parts.pop("xl/featurePropertyBag/featurePropertyBag.xml")
    )
    with pytest.raises(CleaningError, match="featurePropertyBag"):
        cleaner.analyze(feature_bag_xlsx, CleaningPlan())


def test_xlsx_failed_publication_leaves_no_partial_file(cleaner, tmp_path, monkeypatch):
    path = tmp_path / "source.xlsx"
    make_book(path)
    before = fingerprint(path)
    job = cleaner.analyze(path, CleaningPlan(True, True, True, True))
    real_write = ZipFile.writestr

    def fail_on_sheet(archive, info, payload, *args, **kwargs):
        if info.filename == job.source.sheet_part:
            raise OSError("simulated XLSX write failure")
        return real_write(archive, info, payload, *args, **kwargs)

    monkeypatch.setattr(ZipFile, "writestr", fail_on_sheet)
    output = tmp_path / "out.xlsx"
    with pytest.raises(CleaningError, match="simulated XLSX write failure"):
        cleaner.export(job, output)
    assert not output.exists()
    assert not list(tmp_path.glob(".cleaner-*"))
    assert fingerprint(path) == before


def remove_feature_package_declarations(parts):
    """Simulate incomplete producers while retaining the real style-side reference."""
    parts.pop("xl/featurePropertyBag/featurePropertyBag.xml")
    for name, attribute, value in (
        ("xl/_rels/workbook.xml.rels", "Id", "rIdFeatureBag"),
        (
            "[Content_Types].xml",
            "ContentType",
            "application/vnd.ms-excel.featurepropertybag+xml",
        ),
    ):
        root = ET.fromstring(parts[name])
        for node in list(root):
            if node.get(attribute) == value:
                root.remove(node)
        parts[name] = ET.tostring(root)


@pytest.mark.parametrize("reference", ["xf", "dxf", "uri", "content_type", "unknown"])
def test_orphan_bag_metadata_cannot_bypass_refusal(cleaner, feature_bag_xlsx, reference):
    def change(parts):
        content_types = parts["[Content_Types].xml"]
        remove_feature_package_declarations(parts)
        styles = ET.fromstring(parts["xl/styles.xml"])
        extension = styles.find("s:cellXfs/s:xf/s:extLst/s:ext", NS)
        if reference == "xf":
            extension.set("uri", "{UNKNOWN}")  # Namespace detection is independent of URI.
        elif reference == "dxf":
            complement = extension[0]
            complement.tag = complement.tag.replace("xfComplement", "DXFComplement")
            extension.set("uri", "{0417FA29-78FA-4A13-93AC-8FF0FAFDF519}")
            xf = styles.find("s:cellXfs", NS)[1]
            ext_list = xf.find("s:extLst", NS)
            xf.remove(ext_list)
            main = NS["s"]
            dxfs = ET.SubElement(styles, f"{{{main}}}dxfs", count="1")
            ET.SubElement(dxfs, f"{{{main}}}dxf").append(ext_list)
        elif reference == "uri":
            extension.remove(extension[0])
        elif reference == "content_type":
            parts["[Content_Types].xml"] = content_types
            xf = styles.find("s:cellXfs", NS)[1]
            xf.remove(xf.find("s:extLst", NS))
        elif reference == "unknown":
            extension.set("uri", "{UNKNOWN}")
            extension[0].tag = "{urn:unknown}UnknownFeature"
        parts["xl/styles.xml"] = ET.tostring(styles)

    rewrite(feature_bag_xlsx, change)
    before = fingerprint(feature_bag_xlsx)
    error = "расширения стилей" if reference == "unknown" else "featurePropertyBag"
    with pytest.raises(CleaningError, match=error):
        cleaner.analyze(feature_bag_xlsx, CleaningPlan(True, True, True, True))
    assert fingerprint(feature_bag_xlsx) == before


@pytest.mark.parametrize("kind", ["rel", "bag_extension", "default", "unknown_attribute"])
def test_additional_feature_bag_semantics_are_rejected(cleaner, feature_bag_xlsx, kind):
    def change(parts):
        name = "xl/featurePropertyBag/featurePropertyBag.xml"
        root = ET.fromstring(parts[name])
        namespace = root.tag.split("}")[0] + "}"
        checkbox = root[0]
        if kind == "rel":
            ET.SubElement(checkbox, namespace + "rel", k="UnknownLink").text = "rIdMissing"
        elif kind == "bag_extension":
            root.insert(0, ET.Element(namespace + "bagExt"))
            checkbox.set("bagExtId", "0")
        elif kind == "default":
            ET.SubElement(checkbox, namespace + "i", k="default").text = "2"
            sheet = ET.fromstring(parts["xl/worksheets/sheet1.xml"])
            cell = sheet.find('s:sheetData/s:row/s:c[@r="A2"]', NS)
            cell.remove(cell.find("s:v", NS))
            cell.attrib.pop("t")  # Empty value does not imply absence of a control.
            parts["xl/worksheets/sheet1.xml"] = ET.tostring(sheet)
        else:
            checkbox.set("unknown", "1")
        parts[name] = ET.tostring(root)

    rewrite(feature_bag_xlsx, change)
    before = fingerprint(feature_bag_xlsx)
    with pytest.raises(CleaningError, match="featurePropertyBag"):
        cleaner.analyze(feature_bag_xlsx, CleaningPlan(remove_empty_rows=True))
    assert fingerprint(feature_bag_xlsx) == before


def test_openpyxl_read_save_loses_real_feature_metadata(feature_bag_xlsx, tmp_path):
    """Diagnostic regression: value readability is not evidence of safe export."""
    before = fingerprint(feature_bag_xlsx)
    bag_part = "xl/featurePropertyBag/featurePropertyBag.xml"
    namespace = "http://schemas.microsoft.com/office/spreadsheetml/2022/featurepropertybag"
    with ZipFile(feature_bag_xlsx) as original:
        bag = ET.fromstring(original.read(bag_part))
        assert [node.get("type") for node in bag] == [
            "Checkbox",
            "XFControls",
            "XFComplement",
            "XFComplements",
        ]
        styles = ET.fromstring(original.read("xl/styles.xml"))
        assert styles.find(f".//{{{namespace}}}xfComplement").get("i") == "0"
        for name in ("[Content_Types].xml", "xl/_rels/workbook.xml.rels"):
            assert b"featurepropertybag" in original.read(name).lower()
    # Both reading modes expose the boolean but not the control's semantics.
    for read_only in (True, False):
        book = load_workbook(feature_bag_xlsx, read_only=read_only)
        try:
            assert book.active["A2"].value is True
            assert book.active["A2"].data_type == "b"
            if not read_only:
                book.save(tmp_path / "openpyxl-resaved.xlsx")
        finally:
            book.close()
    with ZipFile(tmp_path / "openpyxl-resaved.xlsx") as saved:
        assert saved.testzip() is None
        assert bag_part not in saved.namelist()
        for name in ("[Content_Types].xml", "xl/_rels/workbook.xml.rels", "xl/styles.xml"):
            assert b"featurepropertybag" not in saved.read(name).lower()
        assert b"xfComplement" not in saved.read("xl/styles.xml")
    assert fingerprint(feature_bag_xlsx) == before


def test_empty_unreferenced_bag_is_not_implicitly_allowed(cleaner, feature_bag_xlsx):
    def change(parts):
        name = "xl/featurePropertyBag/featurePropertyBag.xml"
        root = ET.fromstring(parts[name])
        root.clear()
        root.set("count", "0")
        parts[name] = ET.tostring(root)
        styles = ET.fromstring(parts["xl/styles.xml"])
        xf = styles.find("s:cellXfs", NS)[1]
        xf.remove(xf.find("s:extLst", NS))
        parts["xl/styles.xml"] = ET.tostring(styles)

    rewrite(feature_bag_xlsx, change)
    before = fingerprint(feature_bag_xlsx)
    with pytest.raises(CleaningError, match="featurePropertyBag"):
        cleaner.analyze(feature_bag_xlsx, CleaningPlan())
    assert fingerprint(feature_bag_xlsx) == before
