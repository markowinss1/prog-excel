"""Logical row counts are independent of worksheet dimensions and physical tails."""

from xml.etree import ElementTree as ET
from zipfile import ZipFile

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

from excel_toolkit.application.cleaning import CleaningService
from excel_toolkit.infrastructure.files.cleaning import FileCleaningStorage
from excel_toolkit.infrastructure.files.cleaning_xlsx import MAIN, NS
from excel_toolkit.modules.cleaner import CleaningError, CleaningPlan


@pytest.fixture
def cleaner():
    return CleaningService(FileCleaningStorage())


def sheet_xml(path):
    with ZipFile(path) as archive:
        assert archive.testzip() is None
        return ET.fromstring(archive.read("xl/worksheets/sheet1.xml"))


def change_sheet(path, change):
    with ZipFile(path) as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}
    name = "xl/worksheets/sheet1.xml"
    root = ET.fromstring(parts[name])
    change(root)
    parts[name] = ET.tostring(root)
    with ZipFile(path, "w") as archive:
        for name, value in parts.items():
            archive.writestr(name, value)


@pytest.mark.parametrize("freeze", [None, "A2"])
@pytest.mark.parametrize("tail", ["row", "cell", "style", "height", "space", "inline", "value"])
def test_multiple_trailing_rows_have_explicit_scope(cleaner, tmp_path, freeze, tail):
    path = tmp_path / "input.xlsx"
    book = Workbook()
    sheet = book.active
    for row in (["id"], ["001"], [None], ["002"], [None], [None]):
        sheet.append(row)
    sheet.freeze_panes = freeze
    # Formatting does not protect an empty row *inside* the logical rectangle.
    sheet["A3"].font = Font(bold=True)
    for index in (5, 6):
        if tail == "cell":
            sheet.cell(index, 1, "")
        elif tail == "style":
            sheet.cell(index, 1).font = Font(italic=True)
        elif tail == "height":
            sheet.row_dimensions[index].height = 30
        elif tail == "space":
            sheet.cell(index, 1, " \t ")
    book.create_sheet("Other").append([" untouched ", "00009"])
    book.save(path)
    book.close()
    if tail in {"inline", "value"}:

        def payloads(root):
            for index in (5, 6):
                row = root.find(f's:sheetData/s:row[@r="{index}"]', NS)
                cell = ET.SubElement(row, f"{{{MAIN}}}c", r=f"A{index}")
                if tail == "inline":
                    cell.set("t", "inlineStr")
                    ET.SubElement(ET.SubElement(cell, f"{{{MAIN}}}is"), f"{{{MAIN}}}t")
                else:
                    ET.SubElement(cell, f"{{{MAIN}}}v")

        change_sheet(path, payloads)
    before = path.read_bytes(), path.stat().st_mtime_ns
    # max_row differs even though all variants contain six physical row elements.
    assert len(sheet_xml(path).findall("s:sheetData/s:row", NS)) == 6
    loaded = load_workbook(path)
    assert loaded.active.max_row == (4 if tail in {"row", "height"} else 6)
    loaded.close()
    job = cleaner.analyze(path, CleaningPlan(True, True, True, True))
    explicit_payload = tail in {"space", "inline", "value"}
    report = job.result.report
    assert report.original_rows == (6 if explicit_payload else 4)
    assert report.removed_empty_rows == (3 if explicit_payload else 1)
    assert report.final_rows == 3
    assert job.source.xlsx_trailing_rows == (0 if explicit_payload else 2)
    assert bool(report.warnings) is not explicit_payload
    if not explicit_payload:
        assert "XML-строки: 2" in report.warnings[0]
    output = tmp_path / "out.xlsx"
    cleaner.export(job, output)
    assert (path.read_bytes(), path.stat().st_mtime_ns) == before
    rows = sheet_xml(output).findall("s:sheetData/s:row", NS)
    assert [int(row.get("r")) for row in rows] == list(range(1, 4 if explicit_payload else 6))
    assert sheet_xml(output).find("s:dimension", NS) is None
    with ZipFile(path) as original, ZipFile(output) as result:
        for name in original.namelist():
            if name != job.source.sheet_part:
                assert original.read(name) == result.read(name)
    loaded = load_workbook(output)
    try:
        assert loaded.active["A2"].value == "001"
        assert loaded.active["A2"].data_type == "s"
        assert loaded.active.freeze_panes == freeze
        assert loaded["Other"]["B1"].value == "00009"
        if tail == "style":
            assert loaded.active["A4"].font.italic
            assert loaded.active["A5"].font.italic
        if tail == "height":
            assert loaded.active.row_dimensions[4].height == 30
            assert loaded.active.row_dimensions[5].height == 30
    finally:
        loaded.close()
    again = cleaner.analyze(output, CleaningPlan(True, True, True, True))
    assert again.source.table == job.result.table
    assert again.result.report.original_rows == report.final_rows
    assert again.result.report.removed_empty_rows == 0


@pytest.mark.parametrize("has_header", [False, True])
def test_leading_and_interior_blank_rows(cleaner, tmp_path, has_header):
    path = tmp_path / "leading.xlsx"
    book = Workbook()
    for row in ([None], ["001"], [None], ["002"]):
        book.active.append(row)
    selection = book.active.sheet_view.selection[0]
    selection.activeCell = selection.sqref = "A2"
    book.save(path)
    book.close()
    job = cleaner.analyze(path, CleaningPlan(remove_empty_rows=True, has_header=has_header))
    assert job.result.report.original_rows == 4
    assert job.result.report.removed_empty_rows == (1 if has_header else 2)
    assert job.result.source_rows == ((1, 2, 4) if has_header else (2, 4))
    output = tmp_path / "out.xlsx"
    cleaner.export(job, output)
    assert cleaner.analyze(output, CleaningPlan()).source.table == job.result.table
    assert len(sheet_xml(output).findall("s:sheetData/s:row", NS)) == len(job.result.source_rows)


@pytest.mark.parametrize("header", [False, True])
@pytest.mark.parametrize("tail", [False, True])
def test_header_only_or_empty_sheet_and_inflated_dimension(cleaner, tmp_path, header, tail):
    path = tmp_path / "header.xlsx"
    book = Workbook()
    if header:
        book.active.append(["id"])
    if tail:
        book.active.row_dimensions[9].height = 32
        book.active["B10"].font = Font(bold=True)
    book.save(path)
    book.close()
    change_sheet(path, lambda root: root.find("s:dimension", NS).set("ref", "A1:XFD1048576"))
    job = cleaner.analyze(path, CleaningPlan(True, True, True, True))
    assert job.result.report.original_rows == job.result.report.final_rows == int(header)
    assert job.result.report.removed_empty_rows == 0
    assert job.source.xlsx_trailing_rows == (2 if tail else 0)
    output = tmp_path / "out.xlsx"
    cleaner.export(job, output)
    # Without changes even a stale dimension is retained byte-for-byte; it was not trusted.
    with ZipFile(path) as original, ZipFile(output) as result:
        for name in original.namelist():
            assert original.read(name) == result.read(name)
    assert cleaner.analyze(output, CleaningPlan()).source.table == job.result.table


def test_trim_without_removal_keeps_blank_payload_and_report(cleaner, tmp_path):
    path = tmp_path / "trim.xlsx"
    book = Workbook()
    for row in (["id"], ["001"], ["  "], [None]):
        book.active.append(row)
    book.save(path)
    book.close()
    job = cleaner.analyze(path, CleaningPlan(trim_whitespace=True))
    assert job.result.report.original_rows == job.result.report.final_rows == 3
    assert job.result.report.removed_empty_rows == 0
    output = tmp_path / "out.xlsx"
    cleaner.export(job, output)
    root = sheet_xml(output)
    assert root.find('s:sheetData/s:row[@r="3"]/s:c/s:is/s:t', NS) is not None
    assert root.find('s:sheetData/s:row[@r="4"]', NS) is not None
    assert cleaner.analyze(output, CleaningPlan()).result.report.original_rows == 3


def test_formula_on_other_sheet_blocks_blank_row_deletion(cleaner, tmp_path):
    path = tmp_path / "formula.xlsx"
    book = Workbook()
    for row in (["id"], ["001"], [None], ["002"], [None]):
        book.active.append(row)
    book.create_sheet("Formula")["A1"] = "=Sheet!A4"
    book.save(path)
    book.close()
    before = path.read_bytes()
    with pytest.raises(CleaningError, match="формулы"):
        cleaner.analyze(path, CleaningPlan(remove_empty_rows=True))
    job = cleaner.analyze(path, CleaningPlan(trim_whitespace=True))
    assert any("XML-строки: 1" in warning for warning in job.result.report.warnings)
    assert any("кэш" in warning for warning in job.result.report.warnings)
    output = tmp_path / "out.xlsx"
    cleaner.export(job, output)
    assert path.read_bytes() == before
    assert cleaner.analyze(output, CleaningPlan()).result.report.original_rows == 4


@pytest.mark.parametrize("freeze", [None, "A2"])
@pytest.mark.parametrize("tail", [None, " "])
def test_ten_thousand_clients_reproduces_row_discrepancy(cleaner, tmp_path, freeze, tail):
    path = tmp_path / "clients.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.append(["ID", "Name"])
    for index in range(10_000):
        sheet.append([f"{index:06d}", f"Name {index}"])
    for _ in range(199):
        sheet.append([None])
    for index in range(400):
        sheet.append([f"{index:06d}", f"Name {index}"])
    sheet.append([tail])
    sheet.freeze_panes = freeze
    assert sheet.max_row == 10_601
    book.save(path)
    book.close()
    before = path.read_bytes(), path.stat().st_mtime_ns
    root = sheet_xml(path)
    last = root.find('s:sheetData/s:row[@r="10601"]', NS)
    assert last is not None
    assert len(last) == (0 if tail is None else 1)
    assert root.find("s:dimension", NS).get("ref") == "A1:B10601"
    job = cleaner.analyze(path, CleaningPlan(True, True, True, True))
    report = job.result.report
    assert report.original_rows == (10_600 if tail is None else 10_601)
    assert report.removed_empty_rows == (199 if tail is None else 200)
    assert report.removed_duplicates == 400
    assert report.final_rows == 10_001
    assert report.original_rows - report.removed_empty_rows - report.removed_duplicates == (
        report.final_rows
    )
    output = tmp_path / "out.xlsx"
    cleaner.export(job, output)
    rows = sheet_xml(output).findall("s:sheetData/s:row", NS)
    assert len(rows) == (10_002 if tail is None else 10_001)
    if tail is None:
        assert rows[-1].attrib == {"r": "10002"} and len(rows[-1]) == 0
    again = cleaner.analyze(output, CleaningPlan(True, True, True, True))
    assert again.source.table == job.result.table
    assert again.result.report.original_rows == 10_001
    assert again.result.report.removed_empty_rows == again.result.report.removed_duplicates == 0
    assert (path.read_bytes(), path.stat().st_mtime_ns) == before


def test_protected_empty_header_is_physically_retained(cleaner, tmp_path):
    path = tmp_path / "blank-header.xlsx"
    book = Workbook()
    book.active.append([None])
    book.active.append([" "])
    book.save(path)
    book.close()
    job = cleaner.analyze(path, CleaningPlan(remove_empty_rows=True))
    assert job.result.report.original_rows == 2
    assert job.result.report.removed_empty_rows == 1
    assert job.result.report.final_rows == 1
    assert job.result.source_rows == (1,)
    output = tmp_path / "out.xlsx"
    cleaner.export(job, output)
    rows = sheet_xml(output).findall("s:sheetData/s:row", NS)
    assert len(rows) == 1 and rows[0].get("r") == "1" and len(rows[0]) == 0
    # A fresh read determines its own rectangle: no payload now anchors a table.
    again = cleaner.analyze(output, CleaningPlan())
    assert again.result.report.original_rows == 0
    assert again.source.xlsx_trailing_rows == 1
