"""Advanced cleaning through the real file adapters and guarded export service."""

import csv
import hashlib
from zipfile import ZipFile

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

from excel_toolkit.application.cleaning import CleaningService
from excel_toolkit.application.models import CsvOptions
from excel_toolkit.infrastructure.files.cleaning import FileCleaningStorage
from excel_toolkit.modules.cleaner import CleaningError, CleaningPlan


def fingerprint(path):
    return hashlib.sha256(path.read_bytes()).digest(), path.stat().st_mtime_ns


def make_source(path, *, freeze=None):
    rows = [
        ["Client ID", "Name", "City", "Email"],
        [" 00123 ", " Alex ", "Astana", ""],
        ["00123", "Alex", "Almaty", "alex@example.invalid"],
        ["", "Unknown", "Astana", ""],
        ["", "Unknown", "Astana", ""],
        ["", "", "", ""],
        ["00007", "Dana", "Astana", ""],
        ["00007", "Dana", "Astana", ""],
        ["abc", "Lower", "Astana", "lower@example.invalid"],
        ["ABC", "Upper", "Astana", "upper@example.invalid"],
    ]
    if path.suffix == ".csv":
        with path.open("x", encoding="utf-8-sig", newline="") as stream:
            csv.writer(stream, delimiter=";").writerows(rows)
    else:
        book = Workbook()
        sheet = book.active
        sheet.title = "Clients"
        for row in rows:
            sheet.append(row)
        sheet.freeze_panes = freeze
        sheet["A7"].font = Font(bold=True)
        sheet.row_dimensions[7].height = 31
        book.create_sheet("Other").append([" unchanged ", "00009"])
        book.save(path)
        book.close()


@pytest.mark.parametrize("suffix,freeze", [(".csv", None), (".xlsx", None), (".xlsx", "A2")])
def test_selected_export_requires_consent_preserves_source_types_and_other_parts(
    tmp_path, suffix, freeze
):
    source = tmp_path / f"source{suffix}"
    make_source(source, freeze=freeze)
    before = fingerprint(source)
    service = CleaningService(FileCleaningStorage())
    plan = CleaningPlan(True, True, True, True, duplicate_columns=(0,))
    job = service.analyze(source, plan, csv_options=CsvOptions(delimiter=";"))
    report = job.result.report
    assert (report.original_rows, report.final_rows, report.removed_empty_rows) == (10, 7, 1)
    assert report.removed_duplicates == 2
    assert report.duplicate_group_count == 2
    assert report.conflicting_duplicate_groups == 1
    assert report.missing_key_rows == 2
    assert report.missing_values == report.issue_count == 6
    assert job.result.source_rows == (1, 2, 4, 5, 7, 9, 10)
    assert [(i.source_row, i.column, i.kind) for i in report.issues] == [
        (2, 4, "Missing Value"),
        (4, 1, "Missing Key"),
        (4, 4, "Missing Value"),
        (5, 1, "Missing Key"),
        (5, 4, "Missing Value"),
        (7, 4, "Missing Value"),
    ]
    assert report.duplicate_groups[0].kept_source_row == 2
    assert report.duplicate_groups[0].removed_source_rows == (3,)
    assert report.duplicate_groups[0].conflicting_columns == (2, 3)
    destination = tmp_path / f"cleaned{suffix}"
    with pytest.raises(CleaningError, match="подтверждения.*1"):
        service.export(job, destination)
    assert not destination.exists()
    service.export(job, destination, confirm_conflicts=True)
    assert fingerprint(source) == before
    assert not list(tmp_path.glob(".cleaner-*"))
    # Consent is per invocation, never stored on the reusable job or service.
    with pytest.raises(CleaningError, match="подтверждения"):
        service.export(job, tmp_path / f"another{suffix}")
    with pytest.raises(CleaningError, match="оригинала"):
        service.export(job, source, confirm_conflicts=True)
    with pytest.raises(CleaningError, match="существует"):
        service.export(job, destination, confirm_conflicts=True)
    if suffix == ".csv":
        with destination.open(encoding="utf-8-sig", newline="") as stream:
            data = list(csv.reader(stream, delimiter=";"))
        assert [row[0] for row in data[1:]] == ["00123", "", "", "00007", "abc", "ABC"]
    else:
        with ZipFile(source) as original, ZipFile(destination) as exported:
            assert exported.testzip() is None
            for part in original.namelist():
                if part != job.source.sheet_part:
                    assert exported.read(part) == original.read(part)
        book = load_workbook(destination)
        try:
            sheet = book["Clients"]
            assert sheet.freeze_panes == freeze
            assert sheet["A2"].value == "00123" and sheet["A2"].data_type == "s"
            assert sheet["A5"].value == "00007" and sheet["A5"].font.bold
            assert sheet.row_dimensions[5].height == 31
            assert sheet["A3"].value is None and sheet["A4"].value is None
            assert book["Other"]["A1"].value == " unchanged "
        finally:
            book.close()
    again = service.analyze(destination, plan, csv_options=CsvOptions(delimiter=";"))
    assert again.result.report.removed_duplicates == 0
    assert again.result.report.missing_key_rows == 2


@pytest.mark.parametrize("consent", [None, 1, "True", "False"])
def test_export_does_not_accept_truthy_values_as_confirmation(tmp_path, consent):
    source = tmp_path / "source.csv"
    make_source(source)
    service = CleaningService(FileCleaningStorage())
    job = service.analyze(
        source,
        CleaningPlan(trim_whitespace=True, remove_duplicates=True, duplicate_columns=(0,)),
        csv_options=CsvOptions(delimiter=";"),
    )
    destination = tmp_path / "blocked.csv"
    with pytest.raises(CleaningError, match="bool"):
        service.export(job, destination, confirm_conflicts=consent)
    assert not destination.exists()


def test_full_conflict_count_guards_export_beyond_report_limit(tmp_path):
    path = tmp_path / "many.csv"
    with path.open("x", encoding="utf-8", newline="") as stream:
        writer = csv.writer(stream)
        writer.writerow(["ID", "City"])
        for index in range(105):
            writer.writerow([f"{index:05d}", "Astana"])
            writer.writerow([f"{index:05d}", "Almaty"])
    service = CleaningService(FileCleaningStorage())
    job = service.analyze(path, CleaningPlan(remove_duplicates=True, duplicate_columns=(0,)))
    assert len(job.result.report.duplicate_groups) == 100
    assert job.result.report.conflicting_duplicate_groups == 105
    with pytest.raises(CleaningError, match="105"):
        service.export(job, tmp_path / "blocked.csv")
    assert not (tmp_path / "blocked.csv").exists()


def test_selected_columns_still_reject_feature_property_bags(tmp_path, feature_bag_xlsx):
    before = fingerprint(feature_bag_xlsx)
    service = CleaningService(FileCleaningStorage())
    with pytest.raises(CleaningError, match="featurePropertyBag"):
        service.analyze(
            feature_bag_xlsx, CleaningPlan(remove_duplicates=True, duplicate_columns=(0,))
        )
    assert fingerprint(feature_bag_xlsx) == before


@pytest.mark.parametrize("structure", ["merge", "formula_other_sheet", "split"])
def test_selected_deletion_keeps_xlsx_restrictions(tmp_path, structure):
    source = tmp_path / "unsafe.xlsx"
    make_source(source, freeze="A2")
    book = load_workbook(source)
    if structure == "merge":
        book["Other"].merge_cells("A1:B1")
    elif structure == "formula_other_sheet":
        book["Other"]["C1"] = "=1+1"
    else:
        book.active.sheet_view.pane.state = "split"
    book.save(source)
    book.close()
    before = fingerprint(source)
    service = CleaningService(FileCleaningStorage())
    with pytest.raises(CleaningError):
        service.analyze(source, CleaningPlan(True, True, True, True, duplicate_columns=(0,)))
    assert fingerprint(source) == before


@pytest.mark.parametrize("suffix", [".csv", ".xlsx"])
def test_ten_thousand_records_selected_analysis_and_full_export(tmp_path, suffix):
    path = tmp_path / f"large{suffix}"
    rows = [["ID", "Name", "City"]]
    rows.extend(
        [f"{i // 2:06d}", f"Client {i // 2}", "Astana" if i % 2 == 0 else "Almaty"]
        for i in range(9900)
    )
    rows.extend([["", "Review", "Astana"] for _ in range(100)])
    if suffix == ".csv":
        with path.open("x", encoding="utf-8-sig", newline="") as stream:
            csv.writer(stream).writerows(rows)
    else:
        book = Workbook()
        for row in rows:
            book.active.append(row)
        book.active.freeze_panes = "A2"
        book.save(path)
        book.close()
    before = fingerprint(path)
    service = CleaningService(FileCleaningStorage())
    job = service.analyze(path, CleaningPlan(True, True, True, True, duplicate_columns=(0,)))
    report = job.result.report
    assert report.original_rows == 10_001
    assert report.final_rows == 5051
    assert report.removed_duplicates == report.conflicting_duplicate_groups == 4950
    assert report.missing_key_rows == report.issue_count == 100
    assert len(report.duplicate_groups) == 100
    assert len(report.issues) == 100 and not report.issues_truncated
    assert job.result.source_rows[-100:] == tuple(range(9902, 10_002))
    output = tmp_path / f"large_cleaned{suffix}"
    service.export(job, output, confirm_conflicts=True)
    restored = service.analyze(output, CleaningPlan())
    assert restored.result.table == job.result.table
    assert fingerprint(path) == before
