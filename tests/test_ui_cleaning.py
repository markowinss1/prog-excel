import csv
import hashlib
import os
import time
from dataclasses import replace
from pathlib import Path
from threading import Event

import pytest
from openpyxl import Workbook, load_workbook
from PySide6.QtCore import QThread, QTimer
from PySide6.QtWidgets import QApplication, QFileDialog

from excel_toolkit.application.cleaning import CleaningService
from excel_toolkit.application.models import CsvOptions
from excel_toolkit.infrastructure.files.cleaning import FileCleaningStorage
from excel_toolkit.modules.cleaner import CleaningError, CleaningPlan
from excel_toolkit.ui.main_window import MainWindow


@pytest.fixture(scope="module")
def app():
    os.environ["QT_QPA_PLATFORM"] = "offscreen"
    os.environ["QT_QPA_PLATFORMTHEME"] = ""
    return QApplication.instance() or QApplication([])


def wait_until(app, predicate):
    deadline = time.monotonic() + 8
    while not predicate():
        app.processEvents()
        if time.monotonic() > deadline:
            pytest.fail("Timed out waiting for Cleaner background operation")
        time.sleep(0.002)
    app.processEvents()


@pytest.fixture
def window(app, service):
    widget = MainWindow(service, CleaningService(FileCleaningStorage()))
    widget.show()
    yield widget
    widget.close()
    wait_until(app, lambda: widget._worker is None)
    widget.close()
    widget.deleteLater()
    app.processEvents()


@pytest.fixture
def source_csv(tmp_path):
    path = tmp_path / "клиенты.csv"
    path.write_text("ID,Имя\n 001234 , Анна \n001234,Анна\n ,\n000007,\n", encoding="utf-8-sig")
    return path


def load(app, window, path):
    window.load_file(path)
    wait_until(app, lambda: window._worker is None)
    assert window.errors.toPlainText() == ""


def analyze(app, window):
    window.analyze_button.click()
    wait_until(app, lambda: window._worker is None)


def test_initial_controls_and_save_guard(window, monkeypatch):
    def unexpected_dialog(*args):
        pytest.fail("Save dialog must not open before successful analysis")

    monkeypatch.setattr(QFileDialog, "getSaveFileName", unexpected_dialog)
    assert all(box.isChecked() for box in window.operation_checks.values())
    assert window.cleaning_plan() == CleaningPlan(True, True, True, True)
    assert not window.analyze_button.isEnabled()
    assert not window.save_button.isEnabled()
    assert not window.preview_tabs.isTabEnabled(1)
    window._choose_destination()
    window.save_cleaned_file(Path("must-not-exist.csv"))
    assert window._worker is None


def test_csv_end_to_end_report_tabs_save_dialog_and_original_unchanged(
    app, window, source_csv, tmp_path, monkeypatch
):
    before = source_csv.read_bytes(), source_csv.stat().st_mtime_ns
    destination = tmp_path / "cleaned.csv"
    monkeypatch.setattr(QFileDialog, "getOpenFileName", lambda *args: (str(source_csv), ""))
    save_calls = []

    def save_dialog(*args, **kwargs):
        save_calls.append(args)
        return str(destination), "CSV (*.csv)"

    monkeypatch.setattr(QFileDialog, "getSaveFileName", save_dialog)
    gui_threads = []
    window.model.modelReset.connect(lambda: gui_threads.append(QThread.currentThread()))
    window.cleaned_model.modelReset.connect(lambda: gui_threads.append(QThread.currentThread()))
    window.open_button.click()
    wait_until(app, lambda: window._worker is None)
    assert window.model.rowCount() == 5
    original_preview = window.model.preview
    assert original_preview.rows[1][0].value == " 001234 "
    assert window.analyze_button.isEnabled()
    assert not window.save_button.isEnabled()
    analyze(app, window)
    assert window.errors.toPlainText() == ""
    assert window.save_button.isEnabled()
    assert window.preview_tabs.currentIndex() == 1
    assert {k: label.text() for k, label in window.summary_values.items()} == {
        "original_rows": "5",
        "final_rows": "3",
        "removed_empty_rows": "1",
        "removed_duplicates": "1",
        "changed_text_cells": "3",
        "missing_values": "1",
    }
    details = window.report_details.toPlainText()
    assert "Trim Whitespace, Remove Empty Rows, Remove Duplicates, Detect Missing Values" in details
    assert window.cleaned_model.data(window.cleaned_model.index(1, 0)) == "001234"
    window.preview_tabs.setCurrentIndex(0)
    assert window.model.data(window.model.index(1, 0)) == " 001234 "
    assert window.model.rowCount() == 5
    window.preview_tabs.setCurrentIndex(1)
    assert window.cleaned_model.rowCount() == 3
    assert original_preview.rows[1][0].value == " 001234 "
    window.save_button.click()
    wait_until(app, lambda: window._worker is None)
    assert save_calls[0][2] == str(source_csv.with_name("клиенты_cleaned.csv"))
    assert save_calls[0][3] == "CSV (*.csv)"
    assert window.saved_label.isVisible()
    assert str(destination.absolute()) in window.saved_label.text()
    with destination.open(encoding="utf-8-sig", newline="") as stream:
        assert list(csv.reader(stream)) == [["ID", "Имя"], ["001234", "Анна"], ["000007", ""]]
    assert before == (source_csv.read_bytes(), source_csv.stat().st_mtime_ns)
    assert all(thread == app.thread() for thread in gui_threads)
    assert window._worker is None
    assert not window.progress.isVisible()


@pytest.mark.parametrize(
    "flags",
    [
        (True, False, False, False),
        (False, True, False, False),
        (False, False, True, False),
        (False, False, False, True),
        (False, False, False, False),
    ],
)
def test_checkbox_plan_reaches_service_and_summary_comes_from_report(
    app, window, source_csv, flags
):
    calls = []

    class RecordingService(CleaningService):
        def analyze(self, path, plan, **kwargs):
            calls.append((plan, kwargs, QThread.currentThread()))
            job = super().analyze(path, plan, **kwargs)
            # Display must use report fields, even when they differ from visible row counts.
            report = replace(
                job.result.report,
                original_rows=12345,
                missing_values=321,
                warnings=("Проверочное предупреждение\nВторая строка",),
            )
            return replace(job, result=replace(job.result, report=report))

    window._cleaning_service = RecordingService(FileCleaningStorage())
    load(app, window, source_csv)
    for check, flag in zip(window.operation_checks.values(), flags, strict=True):
        check.setChecked(flag)
    analyze(app, window)
    assert calls[0][0] == CleaningPlan(*flags)
    assert calls[0][1] == {"sheet": None, "csv_options": CsvOptions()}
    assert calls[0][2] != app.thread()
    assert window.summary_values["original_rows"].text() == "12345"
    assert window.summary_values["missing_values"].text() == "321"
    assert "Проверочное предупреждение\nВторая строка" in window.report_details.toPlainText()
    if not any(flags):
        assert "не выбраны" in window.report_details.toPlainText()
        assert window.model.preview.rows == window.cleaned_model.preview.rows


@pytest.mark.parametrize(
    "change",
    [
        "trim_whitespace",
        "remove_empty_rows",
        "remove_duplicates",
        "detect_missing_values",
        "encoding",
        "delimiter",
        "reload",
        "file",
    ],
)
def test_result_invalidated_by_settings_or_source_changes(
    app, window, source_csv, tmp_path, change
):
    load(app, window, source_csv)
    analyze(app, window)
    assert window.save_button.isEnabled()
    if change in window.operation_checks:
        window.operation_checks[change].setChecked(False)
    elif change == "encoding":
        window.encoding_combo.setCurrentIndex(1)
    elif change == "delimiter":
        window.delimiter_combo.setCurrentIndex(1)
    elif change == "reload":
        window.reload_button.click()
    elif change == "file":
        another = tmp_path / "another.csv"
        another.write_text("h\nx\n", encoding="utf-8")
        window.load_file(another)
    assert not window.save_button.isEnabled()
    assert window._job is None
    assert window.cleaned_model.rowCount() == 0
    assert not window.preview_tabs.isTabEnabled(1)
    assert all(label.text() == "—" for label in window.summary_values.values())
    window.save_cleaned_file(tmp_path / "stale.csv")
    wait_until(app, lambda: window._worker is None)
    assert not (tmp_path / "stale.csv").exists()
    assert not window.save_button.isEnabled()


def test_sheet_selection_invalidates_and_analyzes_selected_xlsx_sheet(
    app, window, tmp_path, monkeypatch
):
    path = tmp_path / "book.xlsx"
    book = Workbook()
    book.active.title = "Первый"
    book.active.append([" h "])
    book.active.append([" one "])
    second = book.create_sheet("Второй")
    second.append(["id"])
    second.append([" 002 "])
    second.append(["002"])
    book.save(path)
    book.close()
    before = hashlib.sha256(path.read_bytes()).digest()
    load(app, window, path)
    analyze(app, window)
    window.sheet_combo.setCurrentText("Второй")
    assert not window.save_button.isEnabled()
    assert window.model.rowCount() == 0
    window.sheet_combo.textActivated.emit("Второй")
    wait_until(app, lambda: window._worker is None)
    assert window.model.rowCount() == 3
    analyze(app, window)
    assert window._job.source.selected_sheet == "Второй"
    assert window.summary_values["removed_duplicates"].text() == "1"
    assert window.cleaned_model.rowCount() == 2
    destination = tmp_path / "out.xlsx"
    monkeypatch.setattr(
        QFileDialog, "getSaveFileName", lambda *args, **kwargs: (str(destination), "")
    )
    window.save_button.click()
    wait_until(app, lambda: window._worker is None)
    result = load_workbook(destination)
    try:
        assert result["Первый"]["A1"].value == " h "
        assert result["Второй"]["A2"].value == "002"
        assert result["Второй"].max_row == 2
    finally:
        result.close()
    assert hashlib.sha256(path.read_bytes()).digest() == before


def test_csv_options_used_for_full_analysis_without_requiring_reload(app, window, tmp_path):
    path = tmp_path / "legacy.csv"
    path.write_text("Код;Имя\n 001 ; Тест \n001;Тест\n", encoding="cp1251")
    window.load_file(path)
    wait_until(app, lambda: window._worker is None)
    assert window.errors.toPlainText()
    window.encoding_combo.setCurrentIndex(window.encoding_combo.findData("cp1251"))
    window.delimiter_combo.setCurrentIndex(window.delimiter_combo.findData(";"))
    analyze(app, window)
    assert window.errors.toPlainText() == ""
    assert window._job.source.csv_options == CsvOptions("cp1251", ";")
    assert window.cleaned_model.rowCount() == 2
    assert window.cleaned_model.data(window.cleaned_model.index(1, 1)) == "Тест"


def test_full_analysis_and_export_are_not_limited_to_preview(app, window, tmp_path):
    path = tmp_path / "large.csv"
    path.write_text(
        "id\n" + "".join(f" {i:05} \n" for i in range(250)) + " 00249 \n", encoding="utf-8"
    )
    load(app, window, path)
    analyze(app, window)
    assert window.model.rowCount() == window.cleaned_model.rowCount() == 100
    assert window.model.preview.more_rows
    assert window.cleaned_model.preview.more_rows
    assert window.summary_values["original_rows"].text() == "252"
    assert window.summary_values["final_rows"].text() == "251"
    assert "за пределами" in window.preview_label.text()
    output = tmp_path / "all.csv"
    window.save_cleaned_file(output)
    wait_until(app, lambda: window._worker is None)
    with output.open(encoding="utf-8-sig", newline="") as stream:
        rows = list(csv.reader(stream))
    assert len(rows) == 251
    assert rows[-1] == ["00249"]


@pytest.mark.parametrize("unsafe", ["formula", "merge"])
def test_safety_refusal_retains_original_and_disables_export(app, window, tmp_path, unsafe):
    path = tmp_path / "unsafe.xlsx"
    book = Workbook()
    book.active.append(["h"])
    book.active.append([" x "])
    book.active.append(["x"])
    if unsafe == "formula":
        book.create_sheet("Formula")["A1"] = "=COUNTA(Sheet!A:A)"
    else:
        book.active.merge_cells("B1:C1")
    book.save(path)
    book.close()
    before = path.read_bytes(), path.stat().st_mtime_ns
    load(app, window, path)
    analyze(app, window)
    assert (
        "формулы" in window.errors.toPlainText()
        if unsafe == "formula"
        else ("mergeCells" in window.errors.toPlainText())
    )
    assert not window.save_button.isEnabled()
    assert window._job is None
    assert window.model.rowCount() == 3
    assert before == (path.read_bytes(), path.stat().st_mtime_ns)
    assert window.analyze_button.isEnabled()
    if unsafe == "formula":
        window.operation_checks["remove_duplicates"].setChecked(False)
        analyze(app, window)
        assert window.save_button.isEnabled()
        assert "кэш" in window.report_details.toPlainText()


@pytest.mark.parametrize("freeze", ["A2", "B1", "C2"])
def test_gui_freeze_analysis_and_export(app, window, tmp_path, freeze):
    path = tmp_path / "frozen.xlsx"
    book = Workbook()
    book.active.append(["ID"])
    book.active.append([" 001 "])
    book.active.append(["001"])
    book.active.freeze_panes = freeze
    book.save(path)
    book.close()
    load(app, window, path)
    analyze(app, window)
    assert window.errors.toPlainText() == ""
    assert window.save_button.isEnabled()
    output = tmp_path / "out.xlsx"
    window.save_cleaned_file(output)
    wait_until(app, lambda: window._worker is None)
    assert window.saved_label.text()
    book = load_workbook(output)
    try:
        assert book.active.freeze_panes == freeze
        assert book.active["A2"].value == "001"
        assert book.active.max_row == 2
    finally:
        book.close()


@pytest.mark.parametrize("unsafe", ["frozen_row", "split", "feature_bag"])
def test_gui_compatibility_refusal_is_visible_and_cannot_export(
    app, window, tmp_path, feature_bag_xlsx, unsafe
):
    path = feature_bag_xlsx
    expected = "featurePropertyBag"
    if unsafe != "feature_bag":
        path = tmp_path / "unsafe-pane.xlsx"
        book = Workbook()
        book.active.append(["ID"])
        book.active.append([" 001 "])
        book.active.append(["001"])
        book.active.freeze_panes = "A4"
        if unsafe == "split":
            book.active.sheet_view.pane.state = "split"
        expected = "закреплённую область" if unsafe == "frozen_row" else "state=frozen"
        book.save(path)
        book.close()
    before = path.read_bytes(), path.stat().st_mtime_ns
    load(app, window, path)
    original_rows = window.model.rowCount()
    analyze(app, window)
    assert expected in window.errors.toPlainText()
    assert window.errors.isVisible()
    assert window.model.rowCount() == original_rows
    assert window._job is None
    assert not window.save_button.isEnabled()
    assert not window.saved_label.text()
    output = tmp_path / "out.xlsx"
    window.save_cleaned_file(output)
    assert not output.exists()
    assert before == (path.read_bytes(), path.stat().st_mtime_ns)
    if unsafe == "frozen_row":
        window.operation_checks["remove_duplicates"].setChecked(False)
        analyze(app, window)
        assert window.errors.toPlainText() == ""
        assert window.save_button.isEnabled()


def test_export_failure_clears_success_and_allows_retry(app, window, source_csv, tmp_path):
    load(app, window, source_csv)
    analyze(app, window)
    destination = tmp_path / "first.csv"
    window.save_cleaned_file(destination)
    wait_until(app, lambda: window._worker is None)
    assert window.saved_label.text()
    for rejected in (source_csv, destination, tmp_path / "wrong.xlsx"):
        window.save_cleaned_file(rejected)
        assert not window.saved_label.text()
        wait_until(app, lambda: window._worker is None)
        assert window.errors.toPlainText()
        assert window.saved_label.text() == ""
        assert "не сохранён" in window.status_label.text()
        assert window.save_button.isEnabled()
    window.save_cleaned_file(tmp_path / "retry.csv")
    wait_until(app, lambda: window._worker is None)
    assert window.errors.toPlainText() == ""
    assert "retry.csv" in window.saved_label.text()


def test_dialog_cancellation_preserves_valid_result(app, window, source_csv, monkeypatch):
    load(app, window, source_csv)
    analyze(app, window)
    job = window._job
    monkeypatch.setattr(QFileDialog, "getOpenFileName", lambda *args: ("", ""))
    monkeypatch.setattr(QFileDialog, "getSaveFileName", lambda *args, **kwargs: ("", ""))
    window.open_button.click()
    window.save_button.click()
    assert window._job is job
    assert window._worker is None
    assert window.save_button.isEnabled()


@pytest.mark.parametrize("operation", ["analyze", "export"])
def test_background_responsive_single_task_and_safe_close(
    app, service, source_csv, tmp_path, operation
):
    started, release = Event(), Event()
    calls, threads = [], []

    class SlowService(CleaningService):
        def analyze(self, *args, **kwargs):
            if operation == "analyze":
                self.pause("analyze")
            return super().analyze(*args, **kwargs)

        def export(self, *args, **kwargs):
            if operation == "export":
                self.pause("export")
            return super().export(*args, **kwargs)

        def pause(self, name):
            calls.append(name)
            threads.append(QThread.currentThread())
            started.set()
            assert release.wait(6), "Test must release background operation"

    widget = MainWindow(service, SlowService(FileCleaningStorage()))
    widget.show()
    ticks = []
    timer = QTimer()
    timer.setInterval(5)
    timer.timeout.connect(lambda: ticks.append(True))
    timer.start()
    try:
        load(app, widget, source_csv)
        if operation == "export":
            analyze(app, widget)
            widget.save_cleaned_file(tmp_path / "out.csv")
        else:
            widget.analyze_button.click()
        wait_until(app, lambda: started.is_set() and len(ticks) >= 3)
        tick_count = len(ticks)
        worker = widget._worker
        assert worker.isRunning()
        assert widget.progress.isVisible()
        assert widget.progress.minimum() == widget.progress.maximum() == 0
        assert not widget.analyze_button.isEnabled()
        assert not widget.save_button.isEnabled()
        assert not widget.open_button.isEnabled()
        assert all(not check.isEnabled() for check in widget.operation_checks.values())
        widget._analyze()
        widget.save_cleaned_file(tmp_path / "duplicate.csv")
        assert widget._worker is worker
        widget.close()
        assert widget.isVisible()
        assert widget._closing
        assert not worker.isInterruptionRequested()  # Cleaner has no fake cancellation.
        wait_until(app, lambda: len(ticks) > tick_count)
        release.set()
        wait_until(app, lambda: widget._worker is None)
        assert not widget.isVisible()
        assert calls == [operation]
        assert all(thread != app.thread() for thread in threads)
        assert not (tmp_path / "duplicate.csv").exists()
        if operation == "export":
            assert (tmp_path / "out.csv").exists()
    finally:
        release.set()
        widget.close()
        wait_until(app, lambda: widget._worker is None)
        widget.close()
        widget.deleteLater()
        timer.stop()
        app.processEvents()


def test_stale_background_analysis_cannot_restore_export(app, window, source_csv):
    started, release = Event(), Event()

    class SlowService(CleaningService):
        def analyze(self, *args, **kwargs):
            started.set()
            assert release.wait(6)
            return super().analyze(*args, **kwargs)

    window._cleaning_service = SlowService(FileCleaningStorage())
    load(app, window, source_csv)
    window.analyze_button.click()
    try:
        wait_until(app, started.is_set)
        # Controls are disabled to users, but programmatic updates must also invalidate the request.
        window.operation_checks["trim_whitespace"].setChecked(False)
    finally:
        release.set()
        wait_until(app, lambda: window._worker is None)
    assert window._job is None
    assert not window.save_button.isEnabled()
    assert window.cleaned_model.rowCount() == 0


@pytest.mark.parametrize(
    "error",
    [CleaningError("Безопасный отказ\n" + "причина " * 400), RuntimeError("unexpected exception")],
)
def test_worker_errors_do_not_crash_or_enable_export(app, window, source_csv, error, caplog):
    class FailingService:
        def analyze(self, *args, **kwargs):
            raise error

    window._cleaning_service = FailingService()
    load(app, window, source_csv)
    analyze(app, window)
    assert window.errors.toPlainText()
    assert not window.save_button.isEnabled()
    assert window.analyze_button.isEnabled()
    if isinstance(error, CleaningError):
        assert window.errors.toPlainText() == str(error)
    else:
        assert "Неожиданная ошибка" in window.errors.toPlainText()
        assert "unexpected exception" in caplog.text


def test_save_dialog_adds_original_extension_and_never_offers_overwrite(
    app, window, source_csv, tmp_path, monkeypatch
):
    load(app, window, source_csv)
    analyze(app, window)

    def save_dialog(*args, **kwargs):
        assert kwargs["options"] & QFileDialog.Option.DontConfirmOverwrite
        assert args[3] == "CSV (*.csv)"
        return str(tmp_path / "without-extension"), "CSV (*.csv)"

    monkeypatch.setattr(QFileDialog, "getSaveFileName", save_dialog)
    window.save_button.click()
    wait_until(app, lambda: window._worker is None)
    assert (tmp_path / "without-extension.csv").exists()
    assert window.csv_hint.isVisible()
    assert "«Текст»" in window.csv_hint.text()


def test_failed_reanalysis_discards_previous_success(app, window, source_csv):
    load(app, window, source_csv)
    analyze(app, window)
    assert window.save_button.isEnabled()
    source_csv.write_text('"unfinished', encoding="utf-8")
    analyze(app, window)
    assert window.errors.toPlainText()
    assert window._job is None
    assert not window.save_button.isEnabled()
    assert window.cleaned_model.rowCount() == 0
    assert window.summary_values["original_rows"].text() == "—"


@pytest.mark.parametrize("size", [(760, 540), (1100, 900)])
def test_small_and_large_layout_keeps_errors_accessible(app, window, source_csv, size):
    window.resize(*size)
    load(app, window, source_csv)

    class FailingService:
        def analyze(self, *args, **kwargs):
            raise CleaningError("Длинная причина отказа: " + "важное сообщение " * 100)

    window._cleaning_service = FailingService()
    analyze(app, window)
    assert window._scroll.horizontalScrollBar().maximum() == 0
    assert window.errors.isVisible()
    viewport = window._scroll.viewport()
    top_left = window.errors.mapTo(viewport, window.errors.rect().topLeft())
    assert viewport.rect().contains(top_left)
    assert "важное сообщение " * 100 in window.errors.toPlainText()
    assert window.errors.isReadOnly()
