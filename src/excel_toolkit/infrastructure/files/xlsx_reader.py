"""Read the first worksheet rows, keeping formulas as source text."""

from pathlib import Path

from openpyxl import load_workbook

from excel_toolkit.application.models import COLUMN_LIMIT, ROW_LIMIT, Cell, Preview, PreviewError
from excel_toolkit.application.preview import CancelCheck, check_cancelled


def read_xlsx(path: Path, sheet: str | None, cancelled: CancelCheck) -> Preview:
    check_cancelled(cancelled)
    # Explicit file ownership also closes the handle if load_workbook itself fails.
    with path.open("rb") as stream:
        workbook = load_workbook(stream, read_only=True, data_only=False, keep_links=False)
        try:
            check_cancelled(cancelled)
            names = tuple(item.title for item in workbook.worksheets)
            if not names:
                raise PreviewError("В книге нет листов с таблицами.")
            selected = sheet if sheet is not None else names[0]
            if selected not in names:
                raise PreviewError(f"Лист «{selected}» не найден.")
            worksheet = workbook[selected]
            # Some producers write incorrect dimensions. Read actual XML rows.
            worksheet.reset_dimensions()
            rows: list[tuple[Cell, ...]] = []
            more_rows = more_columns = False
            iterator = worksheet.iter_rows(max_row=ROW_LIMIT + 1)
            try:
                for index, row in enumerate(iterator):
                    check_cancelled(cancelled)
                    if index == ROW_LIMIT:
                        more_rows = True
                        break
                    more_columns |= len(row) > COLUMN_LIMIT
                    rows.append(
                        tuple(
                            Cell(cell.value, cell.data_type, cell.number_format)
                            for cell in row[:COLUMN_LIMIT]
                        )
                    )
            finally:
                iterator.close()
            return Preview(path, tuple(rows), names, selected, more_rows, more_columns)
        finally:
            workbook.close()
