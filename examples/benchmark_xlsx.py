"""Repeatable warm-cache benchmark; run with the same interpreter on both revisions."""

import argparse
import json
import platform
import statistics
import tempfile
from datetime import datetime
from pathlib import Path
from time import perf_counter

import openpyxl
from openpyxl import Workbook, load_workbook

from excel_toolkit.application.cleaning import CleaningService
from excel_toolkit.infrastructure.files.cleaning import FileCleaningStorage
from excel_toolkit.modules.cleaner import CleaningPlan


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--freeze", default=None, help="Optional freeze coordinate, e.g. A2")
    parser.add_argument("--repeats", type=int, default=5)
    args = parser.parse_args()
    if args.repeats < 1:
        parser.error("--repeats must be positive")
    measurements = []
    with tempfile.TemporaryDirectory(prefix="cleaner-benchmark-") as directory:
        path = Path(directory) / "input.xlsx"
        book = Workbook()
        sheet = book.active
        sheet.append(["ID", "Name", "Number", "Date", "Flag"])
        for index in range(10_000):
            value = index // 2
            sheet.append(
                [None, " "]
                if index % 20 == 19
                else [f" {value:06d} ", f" Name {value} ", value, datetime(2026, 1, 2), True]
            )
        sheet.freeze_panes = args.freeze
        book.create_sheet("Other").append([" untouched ", "00001"])
        book.save(path)
        book.close()
        original = path.read_bytes()
        service = CleaningService(FileCleaningStorage())
        for repetition in range(args.repeats + 1):  # First run is warm-up.
            start = perf_counter()
            job = service.analyze(path, CleaningPlan(True, True, True, True))
            analyzed = perf_counter()
            output = Path(directory) / f"out-{repetition}.xlsx"
            service.export(job, output)
            exported = perf_counter()
            assert job.result.report.original_rows == 10_001
            assert job.result.report.final_rows == 5_001
            assert job.result.report.removed_empty_rows == 500
            assert job.result.report.removed_duplicates == 4_500
            if repetition:
                measurements.append({"analyze": analyzed - start, "export": exported - analyzed})
        assert path.read_bytes() == original
        book = load_workbook(output)
        try:
            assert book.active.max_row == 5_001
            assert book.active["A5001"].value == "004999"
            assert book.active.freeze_panes == args.freeze
            assert book["Other"]["B1"].value == "00001"
        finally:
            book.close()
    print(
        json.dumps(
            {
                "python": platform.python_version(),
                "openpyxl": openpyxl.__version__,
                "platform": platform.platform(),
                "freeze": args.freeze,
                "records": 10_000,
                "samples_seconds": measurements,
                "median_seconds": {
                    key: statistics.median(sample[key] for sample in measurements)
                    for key in ("analyze", "export")
                },
            },
            indent=2,
        )
    )


if __name__ == "__main__":
    main()
